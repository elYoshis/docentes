from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Title
title = doc.add_heading("INFORME DE VERIFICACIÓN DE INSTALACIÓN DE DISTRIBUCIÓN ELÉCTRICA", 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

subtitle = doc.add_paragraph(
    "Áreas: Medicina, Ascensor y Sistema General\n"
    "Universidad Privada Domingo Savio – Sede Oruro"
)
subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# General data
doc.add_paragraph("\nFecha de verificación: [Colocar fecha]")
doc.add_paragraph("Áreas verificadas: Sistema de distribución eléctrica")
doc.add_paragraph("Responsable de verificación: José Manuel Cortez Concha")
doc.add_paragraph("Cargo: Aux. de Mantenimiento y Seguridad")

# Sections
doc.add_heading("1. OBJETIVO DEL INFORME", level=1)
doc.add_paragraph(
    "El presente informe tiene como objetivo documentar la verificación técnica de la instalación del sistema "
    "de distribución eléctrica destinado a las áreas de Medicina, Ascensor y sistema general de la Universidad "
    "Privada Domingo Savio – Sede Oruro, con el fin de constatar que la división de corriente fue realizada "
    "correctamente y conforme a criterios técnicos y de seguridad."
)

doc.add_heading("2. ALCANCE DE LA VERIFICACIÓN", level=1)
for item in [
    "Verificación de la división de la corriente eléctrica para las áreas de Medicina, Ascensor y sistema general.",
    "Revisión del tablero eléctrico y componentes instalados.",
    "Inspección de protecciones eléctricas y switches de control.",
    "Verificación visual del cableado y su correcta identificación."
]:
    doc.add_paragraph(f"- {item}")

doc.add_heading("3. DESCRIPCIÓN DE LA INSTALACIÓN VERIFICADA", level=1)
doc.add_paragraph(
    "La instalación verificada corresponde a un sistema de distribución eléctrica que divide la corriente "
    "eléctrica en circuitos independientes para las áreas de Medicina, Ascensor y sistema general. "
    "Al momento de la verificación, dichos circuitos se encuentran correctamente diferenciados y protegidos, "
    "sin encontrarse aún conectados a las cargas finales de Medicina ni del Ascensor."
)

doc.add_heading("4. MATERIALES UTILIZADOS", level=1)
materials = [
    "Caja de metal para alojamiento del sistema eléctrico.",
    "Panel trasero para sujeción de componentes.",
    "Interruptores termomagnéticos (térmicos).",
    "Switch de encendido y apagado general.",
    "Switch independiente para el área de Medicina.",
    "Switch independiente para el sistema del Ascensor.",
    "Switch independiente para el sistema General.",
    "Cables de cobre de grueso calibre.",
    "Cables de cobre calibre 12."
]
for m in materials:
    doc.add_paragraph(f"- {m}")

doc.add_heading("5. RESULTADOS DE LA VERIFICACIÓN", level=1)
for res in [
    "La corriente eléctrica se encuentra correctamente dividida en circuitos independientes.",
    "Los circuitos de Medicina y Ascensor no se encuentran aún conectados a sus respectivas áreas.",
    "Los interruptores y protecciones instalados se encuentran en buen estado.",
    "No se evidencian cables expuestos ni conexiones sueltas.",
    "La instalación cumple con condiciones de seguridad para su futura conexión."
]:
    doc.add_paragraph(f"- {res}")

doc.add_heading("6. OBSERVACIONES", level=1)
doc.add_paragraph(
    "Se deja constancia que la instalación verificada corresponde únicamente a la división de la corriente "
    "eléctrica. La conexión final a las áreas de Medicina y Ascensor deberá realizarse en una siguiente etapa, "
    "previa verificación técnica correspondiente."
)

doc.add_heading("7. CONCLUSIÓN", level=1)
doc.add_paragraph(
    "De acuerdo con la verificación realizada, se concluye que la instalación del sistema de distribución "
    "eléctrica, correspondiente a la división de corriente para las áreas de Medicina, Ascensor y sistema "
    "general, fue ejecutada correctamente y cumple con las condiciones técnicas y de seguridad necesarias, "
    "quedando apta para su posterior conexión y puesta en servicio."
)

# Signature
doc.add_paragraph("\n\n_______________________________")
doc.add_paragraph("José Manuel Cortez Concha")
doc.add_paragraph("Aux. de Mantenimiento y Seguridad")

# Save file
path = "Informe_Verificacion_Distribucion_Electrica_UPDS_Division.docx"
doc.save(path)

path
